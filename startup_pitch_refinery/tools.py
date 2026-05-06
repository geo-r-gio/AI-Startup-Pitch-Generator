import json
import io
import os
import re
import time
import base64
from contextlib import redirect_stdout
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import quote_plus
from urllib.request import Request, urlopen
from urllib.parse import urlparse

from linkup import LinkupClient
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


class MarketSearchTool:
    """Wrapper around Linkup search with retry/backoff and retrieval diagnostics.

    The adaptive repair controller uses this class in two ways:
    1. broad search for normal market research;
    2. bounded multi-query claim repair search with source quality, domain diversity,
       and claim-snippet overlap diagnostics.
    """

    def __init__(self) -> None:
        self._client = None
        self._max_attempts = self._read_int_env("LINKUP_SEARCH_RETRIES", 2) + 1
        self._retry_backoff_seconds = self._read_float_env(
            "LINKUP_SEARCH_RETRY_BACKOFF_SECONDS",
            0.8,
        )
        self._content_char_limit = self._read_int_env("LINKUP_CONTENT_CHAR_LIMIT", 650)
        self._max_results_total = self._read_int_env("LINKUP_MAX_RESULTS_TOTAL", 8)
        allowed_env = os.getenv("LINKUP_ALLOWED_DOMAINS", "").strip()
        blocked_env = os.getenv("LINKUP_BLOCKED_DOMAINS", "").strip()
        default_blocked = {
            "github.com",
            "marketreportsworld.com",
            "verifiedmarketreports.com",
            "futuredatastats.com",
            "dataintelo.com",
            "marketintelo.com",
            "htfmarketinsights.com",
            "openpr.com",
            "globenewswire.com",
            "einpresswire.com",
        }
        self._allowed_domains = {
            d.strip().lower() for d in allowed_env.split(",") if d.strip()
        }
        blocked_domains = {
            d.strip().lower() for d in blocked_env.split(",") if d.strip()
        }
        self._blocked_domains = blocked_domains or default_blocked
        try:
            self._client = LinkupClient()
        except Exception:
            self._client = None

    @staticmethod
    def _read_int_env(name: str, default: int) -> int:
        try:
            return max(0, int(os.getenv(name, str(default))))
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _read_float_env(name: str, default: float) -> float:
        try:
            return max(0.0, float(os.getenv(name, str(default))))
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _extract_domain(url: str) -> str:
        try:
            parsed = urlparse(url)
            domain = parsed.netloc.lower()
            if domain.startswith("www."):
                domain = domain[4:]
            return domain
        except Exception:
            return ""

    def _is_domain_allowed(self, domain: str) -> bool:
        if not domain:
            return False
        if self._allowed_domains:
            return domain in self._allowed_domains
        return domain not in self._blocked_domains

    @staticmethod
    def _normalize_tokens(text: str) -> List[str]:
        stopwords = {
            "a", "an", "and", "are", "as", "at", "be", "by", "for", "from",
            "has", "have", "in", "into", "is", "it", "its", "of", "on", "or",
            "that", "the", "their", "this", "to", "with", "within", "without",
            "market", "startup", "solution", "solutions", "platform", "users",
        }
        tokens = re.findall(r"[a-zA-Z][a-zA-Z0-9\-]{2,}", (text or "").lower())
        return [tok for tok in tokens if tok not in stopwords]

    @classmethod
    def claim_overlap_score(cls, claim_text: str, candidate_text: str) -> float:
        claim_tokens = set(cls._normalize_tokens(claim_text))
        if not claim_tokens:
            return 0.0
        candidate_tokens = set(cls._normalize_tokens(candidate_text))
        if not candidate_tokens:
            return 0.0
        return round(len(claim_tokens & candidate_tokens) / max(1, len(claim_tokens)), 4)

    @staticmethod
    def source_quality_score(url: str) -> float:
        """Simple source-quality prior used by claim-level repair.

        This is intentionally conservative. It is not a truth oracle; it only gives
        the retry policy a lightweight prior about source credibility.
        """
        url_l = (url or "").lower()
        high_quality = [
            ".gov", ".edu", "who.int", "nih.gov", "cms.gov", "hhs.gov",
            "oecd.org", "worldbank.org", "data.worldbank.org", "imf.org",
            "ec.europa.eu", "eurostat", "mckinsey.com", "deloitte.com",
            "pwc.com", "bcg.com", "bain.com", "cbinsights.com", "rockhealth.com",
            "gartner.com", "forrester.com", "statista.com", "grandviewresearch.com",
            "marketsandmarkets.com", "technavio.com", "imarcgroup.com",
            "fortune.com", "forbes.com", "reuters.com", "apnews.com",
            "businesswire.com", "prnewswire.com",
        ]
        medium_quality = [
            "techcrunch.com", "crunchbase.com", "pitchbook.com", "sifted.eu",
            "builtin.com", "venturebeat.com", "healthcareitnews.com",
        ]
        low_quality = [
            "medium.com", "substack.com", "wikipedia.org", "github.com",
            "marketreportsworld.com", "verifiedmarketreports.com", "futuredatastats.com",
            "dataintelo.com", "marketintelo.com", "htfmarketinsights.com",
            "openpr.com", "einpresswire.com",
        ]
        if any(m in url_l for m in high_quality):
            return 1.0
        if any(m in url_l for m in medium_quality):
            return 0.75
        if any(m in url_l for m in low_quality):
            return 0.35
        return 0.65

    def _compact_result(self, raw: Any) -> Optional[Dict[str, Any]]:
        if isinstance(raw, dict):
            title = raw.get("name", "") or raw.get("title", "")
            url = raw.get("url", "")
            content = raw.get("content", "") or raw.get("snippet", "")
        else:
            title = getattr(raw, "name", "") or getattr(raw, "title", "")
            url = getattr(raw, "url", "")
            content = getattr(raw, "content", "") or getattr(raw, "snippet", "")
        if not url:
            return None
        domain = self._extract_domain(str(url))
        if not self._is_domain_allowed(domain):
            return None
        return {
            "title": str(title or "")[:180],
            "url": str(url),
            "domain": domain,
            "content": str(content or "")[: self._content_char_limit],
            "quality_score": self.source_quality_score(str(url)),
        }

    def _diagnostics_for_results(
        self,
        *,
        results: List[Dict[str, Any]],
        claim_text: str = "",
        dropped_sources: int = 0,
    ) -> Dict[str, Any]:
        sources = [str(item.get("url", "") or "") for item in results if item.get("url")]
        domains = [str(item.get("domain", "") or "") for item in results if item.get("domain")]
        quality_scores = [float(item.get("quality_score", 0.0) or 0.0) for item in results]
        high_quality_source_count = sum(1 for score in quality_scores if score >= 0.65)
        mean_quality_score = round(sum(quality_scores) / max(1, len(quality_scores)), 4) if quality_scores else 0.0
        domain_diversity = round(len(set(domains)) / max(1, len(set(sources))), 4) if sources else 0.0
        overlaps: List[float] = []
        if claim_text:
            for item in results:
                overlaps.append(
                    self.claim_overlap_score(
                        claim_text,
                        f"{item.get('title', '')} {item.get('content', '')}",
                    )
                )
        mean_claim_overlap = round(sum(overlaps) / max(1, len(overlaps)), 4) if overlaps else 0.0
        max_claim_overlap = round(max(overlaps), 4) if overlaps else 0.0
        return {
            "source_count_raw": len(results),
            "unique_source_count": len(set(sources)),
            "high_quality_source_count": high_quality_source_count,
            "mean_quality_score": mean_quality_score,
            "domain_diversity": domain_diversity,
            "mean_claim_overlap": mean_claim_overlap,
            "max_claim_overlap": max_claim_overlap,
            "dropped_sources": dropped_sources,
        }

    def search(self, query: str, claim_text: str = "") -> Dict[str, Any]:
        if self._client is None:
            return {
                "status": "error",
                "query": query,
                "results_json": "[]",
                "results": [],
                "sources": [],
                "dropped_sources": 0,
                "source_count_raw": 0,
                "unique_source_count": 0,
                "high_quality_source_count": 0,
                "mean_quality_score": 0.0,
                "domain_diversity": 0.0,
                "mean_claim_overlap": 0.0,
                "max_claim_overlap": 0.0,
                "attempts": 0,
                "retry_errors": [],
                "error": "Set LINKUP_API_KEY to enable live web search.",
            }

        retry_errors: List[str] = []
        max_attempts = max(1, int(self._max_attempts))
        for attempt in range(1, max_attempts + 1):
            try:
                response = self._client.search(
                    query=query,
                    depth="standard",
                    output_type="searchResults",
                )
                raw_results = getattr(response, "results", [])
                compact: List[Dict[str, Any]] = []
                seen_urls = set()
                dropped_count = 0
                for raw in raw_results:
                    item = self._compact_result(raw)
                    if item is None:
                        dropped_count += 1
                        continue
                    url = item.get("url", "")
                    if not url or url in seen_urls:
                        continue
                    seen_urls.add(url)
                    compact.append(item)

                compact = compact[: max(1, self._max_results_total)]
                sources = [item["url"] for item in compact if item.get("url")]
                diagnostics = self._diagnostics_for_results(
                    results=compact,
                    claim_text=claim_text,
                    dropped_sources=dropped_count,
                )
                return {
                    "status": "ok" if compact else "error",
                    "query": query,
                    "results_json": json.dumps(compact, indent=2),
                    "results": compact,
                    "sources": sources,
                    "attempts": attempt,
                    "retry_errors": retry_errors,
                    "error": "" if compact else "No search results returned.",
                    **diagnostics,
                }
            except Exception as exc:
                retry_errors.append(f"attempt_{attempt}: {exc}")
                if attempt < max_attempts and self._retry_backoff_seconds > 0:
                    time.sleep(self._retry_backoff_seconds * attempt)

        return {
            "status": "error",
            "query": query,
            "results_json": "[]",
            "results": [],
            "sources": [],
            "dropped_sources": 0,
            "source_count_raw": 0,
            "unique_source_count": 0,
            "high_quality_source_count": 0,
            "mean_quality_score": 0.0,
            "domain_diversity": 0.0,
            "mean_claim_overlap": 0.0,
            "max_claim_overlap": 0.0,
            "attempts": max_attempts,
            "retry_errors": retry_errors,
            "error": "Market search failed after " f"{max_attempts} attempt(s): {'; '.join(retry_errors)}",
        }

    def search_many(
        self,
        queries: List[str],
        claim_text: str = "",
        max_results_per_query: int = 3,
    ) -> Dict[str, Any]:
        """Run bounded per-claim searches and deduplicate results by URL.

        `claim_text` is optional for backward compatibility. When provided, the
        return payload includes overlap diagnostics used by the Bayesian retry gate.
        """
        merged: List[Dict[str, Any]] = []
        seen_urls = set()
        statuses: List[Dict[str, Any]] = []
        dropped_sources = 0
        errors: List[str] = []
        total_attempts = 0
        clean_queries = [str(q).strip() for q in queries if str(q).strip()]
        for query in clean_queries:
            payload = self.search(query, claim_text=claim_text)
            total_attempts += int(payload.get("attempts", 0) or 0)
            statuses.append(
                {
                    "query": query,
                    "status": payload.get("status", "unknown"),
                    "source_count": len(payload.get("sources", []) or []),
                    "source_count_raw": payload.get("source_count_raw", 0),
                    "unique_source_count": payload.get("unique_source_count", 0),
                    "high_quality_source_count": payload.get("high_quality_source_count", 0),
                    "mean_claim_overlap": payload.get("mean_claim_overlap", 0.0),
                    "attempts": payload.get("attempts", 0),
                    "retry_errors": payload.get("retry_errors", []),
                    "error": payload.get("error", ""),
                }
            )
            dropped_sources += int(payload.get("dropped_sources", 0) or 0)
            if payload.get("status") != "ok":
                if payload.get("error"):
                    errors.append(str(payload.get("error")))
                continue
            for item in (payload.get("results", []) or [])[:max_results_per_query]:
                if not isinstance(item, dict):
                    continue
                url = str(item.get("url", "") or "")
                if url and url in seen_urls:
                    continue
                if url:
                    seen_urls.add(url)
                item = dict(item)
                item["quality_score"] = self.source_quality_score(url)
                if claim_text:
                    item["claim_overlap"] = self.claim_overlap_score(
                        claim_text,
                        f"{item.get('title', '')} {item.get('content', '')}",
                    )
                else:
                    item.setdefault("claim_overlap", 0.0)
                merged.append(item)

        merged = sorted(
            merged,
            key=lambda item: (
                float(item.get("quality_score", 0.0) or 0.0),
                float(item.get("claim_overlap", 0.0) or 0.0),
            ),
            reverse=True,
        )
        merged = merged[: max(1, self._max_results_total)]
        diagnostics = self._diagnostics_for_results(
            results=merged,
            claim_text=claim_text,
            dropped_sources=dropped_sources,
        )
        return {
            "status": "ok" if merged else "error",
            "query": " | ".join(clean_queries),
            "queries": clean_queries,
            "query_statuses": statuses,
            "results": merged,
            "results_json": json.dumps(merged, indent=2),
            "sources": [item["url"] for item in merged if item.get("url")],
            "attempts": total_attempts,
            "retry_errors": [err for status in statuses for err in (status.get("retry_errors") or [])],
            "error": "; ".join(errors) if errors else ("" if merged else "No search results returned."),
            **diagnostics,
        }


class BusinessCalcTool:
    """Python REPL helper for simple financial projections."""

    def run(self, code: str) -> str:
        try:
            buf = io.StringIO()
            globals_dict = {"__builtins__": {"print": print, "range": range, "len": len}}
            with redirect_stdout(buf):
                exec(code, globals_dict, {})
            return buf.getvalue().strip()
        except Exception as exc:  
            return f"Python calc failed: {exc}"


class GoogleTrendsTool:
    """Google Trends wrapper via pytrends."""

    def __init__(self) -> None:
        self.hl = os.getenv("GOOGLE_TRENDS_HL", "en-US").strip()
        self.tz = int(os.getenv("GOOGLE_TRENDS_TZ", "360").strip() or "360")
        self.geo = os.getenv("GOOGLE_TRENDS_GEO", "US").strip()
        self.timeframe = os.getenv("GOOGLE_TRENDS_TIMEFRAME", "today 12-m").strip()
        self.connect_timeout = int(os.getenv("GOOGLE_TRENDS_CONNECT_TIMEOUT", "6").strip())
        self.read_timeout = int(os.getenv("GOOGLE_TRENDS_READ_TIMEOUT", "12").strip())
        self.retries = int(os.getenv("GOOGLE_TRENDS_RETRIES", "2").strip())
        self.backoff_factor = float(os.getenv("GOOGLE_TRENDS_BACKOFF_FACTOR", "0.4").strip())
        self.max_keywords = int(os.getenv("GOOGLE_TRENDS_MAX_KEYWORDS", "5").strip())
        self._init_error = ""
        try:
            from pytrends.request import TrendReq

            self._client = TrendReq(
                hl=self.hl,
                tz=self.tz,
                timeout=(self.connect_timeout, self.read_timeout),
                retries=self.retries,
                backoff_factor=self.backoff_factor,
            )
        except Exception as exc:  
            self._client = None
            self._init_error = str(exc)

    def fetch(self, keywords: List[str]) -> Dict[str, Any]:
        clean_keywords = [k.strip() for k in keywords if k and k.strip()][: self.max_keywords]
        if not clean_keywords:
            return {
                "status": "error",
                "keywords": [],
                "data": {},
                "error": "No keywords provided for trends lookup.",
            }
        if self._client is None:
            return {
                "status": "error",
                "keywords": clean_keywords,
                "data": {},
                "error": (
                    "Google Trends unavailable. Ensure pytrends is installed in the same "
                    f"Python environment. Details: {self._init_error or 'import failed'}"
                ),
            }

        per_keyword: Dict[str, Any] = {}
        failures: List[str] = []
        for kw in clean_keywords:
            attempts = max(1, self.retries + 1)
            last_error = ""
            for attempt in range(attempts):
                try:
                    self._client.build_payload([kw], timeframe=self.timeframe, geo=self.geo)
                    iot = self._client.interest_over_time()
                    related = self._client.related_queries() or {}

                    summary = {
                        "average_interest": 0.0,
                        "latest_interest": 0,
                        "peak_interest": 0,
                        "top_queries": [],
                        "rising_queries": [],
                    }
                    if not iot.empty and kw in iot.columns:
                        values = [int(v) for v in iot[kw].tolist()]
                        if values:
                            summary["average_interest"] = round(sum(values) / len(values), 2)
                            summary["latest_interest"] = values[-1]
                            summary["peak_interest"] = max(values)

                    kw_related = related.get(kw, {})
                    top_df = kw_related.get("top")
                    if top_df is not None and not top_df.empty and "query" in top_df.columns:
                        summary["top_queries"] = [
                            str(q) for q in top_df["query"].head(5).tolist()
                        ]
                    rising_df = kw_related.get("rising")
                    if (
                        rising_df is not None
                        and not rising_df.empty
                        and "query" in rising_df.columns
                    ):
                        summary["rising_queries"] = [
                            str(q) for q in rising_df["query"].head(5).tolist()
                        ]
                    per_keyword[kw] = summary
                    last_error = ""
                    break
                except Exception as exc:  
                    last_error = str(exc)
                    if attempt < attempts - 1:
                        sleep_s = self.backoff_factor * (2**attempt)
                        time.sleep(max(0.2, sleep_s))
            if last_error:
                failures.append(f"{kw}: {last_error}")

        if not per_keyword:
            return {
                "status": "error",
                "keywords": clean_keywords,
                "data": {},
                "error": "; ".join(failures)[:500] or "Google Trends request failed.",
            }

        status = "ok" if not failures else "partial"
        return {
            "status": status,
            "keywords": clean_keywords,
            "data": per_keyword,
            "error": "; ".join(failures)[:500] if failures else "",
        }


class ScenarioAnalysisTool:
    """Deterministic scenario projections for startup plans."""

    def run(
        self,
        users_year1: int = 8000,
        arpu_monthly: float = 12.0,
        gross_margin: float = 0.72,
    ) -> Dict[str, Any]:
        scenarios = {
            "conservative": {"users_multiplier": 0.7, "arpu_multiplier": 0.9},
            "base": {"users_multiplier": 1.0, "arpu_multiplier": 1.0},
            "aggressive": {"users_multiplier": 1.35, "arpu_multiplier": 1.15},
        }
        output: Dict[str, Any] = {}
        for name, mult in scenarios.items():
            users = int(round(users_year1 * mult["users_multiplier"]))
            arpu = arpu_monthly * mult["arpu_multiplier"]
            revenue = users * arpu * 12
            gross_profit = revenue * gross_margin
            output[name] = {
                "users": users,
                "arpu_monthly": round(arpu, 2),
                "annual_revenue": round(revenue, 2),
                "gross_profit": round(gross_profit, 2),
                "gross_margin": gross_margin,
            }
        return output


def generate_pitch_deck(slides: Dict[str, str], output_path: str) -> str:
    """Create a styled .pptx deck from text content with optional visual enrichment."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    image_dir = path.parent / "images"
    image_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_layout = prs.slide_layouts[0]
    body_layout = prs.slide_layouts[1]

    title_slide = prs.slides.add_slide(title_layout)
    _apply_background(title_slide, RGBColor(17, 24, 39), RGBColor(31, 41, 55))
    title_slide.shapes.title.text = slides.get("title", "Startup Pitch")
    title_slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(44)
    title_slide.shapes.title.text_frame.paragraphs[0].font.bold = True
    title_slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(245, 245, 245)
    title_slide.placeholders[1].text = slides.get("subtitle", "AI Startup Pitch Refinery")
    subtitle_p = title_slide.placeholders[1].text_frame.paragraphs[0]
    subtitle_p.font.size = Pt(22)
    subtitle_p.font.color.rgb = RGBColor(209, 213, 219)

    # Accent stripe for title slide
    stripe = title_slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(6.8), Inches(13.333), Inches(0.7)
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = RGBColor(59, 130, 246)
    stripe.line.fill.background()

    ordered_sections: List[str] = [
        "problem",
        "solution",
        "market",
        "business_model",
        "competitive_advantage",
        "financials",
    ]

    headers = {
        "problem": "Problem",
        "solution": "Solution",
        "market": "Market",
        "business_model": "Business Model",
        "competitive_advantage": "Competitive Advantage",
        "financials": "Financials",
    }

    for idx, key in enumerate(ordered_sections):
        slide = prs.slides.add_slide(body_layout)
        if len(slide.placeholders) > 1:
            ph = slide.placeholders[1]
            ph_el = ph._element
            ph_el.getparent().remove(ph_el)
        section_text = slides.get(key, "TBD")
        accent_color = _section_color(idx)
        _apply_background(slide, RGBColor(249, 250, 251), RGBColor(243, 244, 246))
        _style_section_title(slide, headers[key], accent_color)

        content_box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.5), Inches(7.3), Inches(5.3)
        )
        content_box.fill.solid()
        content_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        content_box.fill.transparency = 0.08
        content_box.line.color.rgb = RGBColor(229, 231, 235)
        content_box.line.width = Pt(1.5)

        _fill_content_text(content_box.text_frame, section_text)

        image_query = f"{headers[key]} startup business"
        image_path = _fetch_slide_image(image_query, image_dir / f"{key}.jpg")

        image_left = Inches(8.3)
        image_top = Inches(1.55)
        image_width = Inches(4.2)
        image_height = Inches(5.2)
        if image_path:
            _add_picture_contain(
                slide,
                str(image_path),
                image_left,
                image_top,
                image_width,
                image_height,
            )
        else:
            fallback = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, image_left, image_top, image_width, image_height
            )
            fallback.fill.solid()
            fallback.fill.fore_color.rgb = RGBColor(224, 231, 255)
            fallback.line.color.rgb = accent_color
            fallback_text = fallback.text_frame
            fallback_text.text = "Visual Placeholder"
            fallback_text.paragraphs[0].font.bold = True
            fallback_text.paragraphs[0].font.size = Pt(18)
            fallback_text.paragraphs[0].font.color.rgb = RGBColor(55, 65, 81)

    prs.save(path)
    return str(path)


def _section_color(idx: int) -> RGBColor:
    palette = [
        RGBColor(37, 99, 235),
        RGBColor(16, 185, 129),
        RGBColor(249, 115, 22),
        RGBColor(168, 85, 247),
        RGBColor(244, 63, 94),
        RGBColor(14, 165, 233),
    ]
    return palette[idx % len(palette)]


def _apply_background(slide, top: RGBColor, bottom: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = top


def _style_section_title(slide, title: str, accent: RGBColor) -> None:
    title_shape = slide.shapes.title
    title_shape.text = title
    p = title_shape.text_frame.paragraphs[0]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 24, 39)

    accent_bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.7), Inches(1.2), Inches(2.2), Inches(0.08)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()


def _fill_content_text(text_frame, content: str) -> None:
    text_frame.clear()
    lines = [ln.strip() for ln in content.split("\n") if ln.strip()]
    if not lines:
        lines = ["TBD"]
    first = text_frame.paragraphs[0]
    first.text = lines[0]
    first.font.size = Pt(20)
    first.font.bold = True
    first.font.color.rgb = RGBColor(17, 24, 39)
    for line in lines[1:]:
        p = text_frame.add_paragraph()
        p.text = line
        p.level = 1
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(55, 65, 81)


def _add_picture_contain(slide, image_path: str, left, top, width, height) -> None:
    """Place image inside target box without distortion and centered."""
    pic = slide.shapes.add_picture(image_path, left, top)
    iw, ih = float(pic.width), float(pic.height)
    bw, bh = float(width), float(height)
    if iw <= 0 or ih <= 0:
        return

    scale = min(bw / iw, bh / ih)
    new_w = int(iw * scale)
    new_h = int(ih * scale)
    pic.width = new_w
    pic.height = new_h
    pic.left = int(float(left) + (bw - new_w) / 2)
    pic.top = int(float(top) + (bh - new_h) / 2)


def _fetch_slide_image(query: str, destination: Path) -> Optional[Path]:
    reuse_cache = os.getenv("IMAGE_REUSE_CACHE", "1").strip().lower() in {"1", "true", "yes", "on"}
    fetch_enabled = os.getenv("IMAGE_FETCH_ENABLED", "1").strip().lower() in {
        "1",
        "true",
        "yes",
        "on",
    }
    if reuse_cache and destination.exists():
        return destination
    if not fetch_enabled:
        return destination if destination.exists() else None

    provider = os.getenv("IMAGE_PROVIDER", "auto").strip().lower()
    providers = [provider]
    if provider == "auto":
        providers = ["unsplash", "openai"]

    for p in providers:
        if p == "openai":
            got = _fetch_from_openai_image(query, destination)
        else:
            got = _fetch_from_unsplash(query, destination)
        if got:
            return got
    return None


def _download_image(url: str, destination: Path, headers: Optional[Dict[str, str]] = None) -> Optional[Path]:
    try:
        req = Request(url, headers=headers or {"User-Agent": "Mozilla/5.0"})
        with urlopen(req, timeout=20) as resp:  # noqa: S310
            content_type = str(resp.headers.get("Content-Type", "")).lower()
            if "image" not in content_type:
                return None
            data = resp.read()
        if not data:
            return None
        destination.write_bytes(data)
        return destination
    except Exception:  
        return None


def _fetch_from_unsplash(query: str, destination: Path) -> Optional[Path]:
    url = f"https://source.unsplash.com/1600x900/?{quote_plus(query)}"
    return _download_image(url, destination)


def _fetch_from_openai_image(query: str, destination: Path) -> Optional[Path]:
    enabled = os.getenv("OPENAI_IMAGE_ENABLED", "0").strip().lower() in {"1", "true", "yes", "on"}
    if not enabled:
        return None
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        return None

    try:
        from openai import OpenAI

        client = OpenAI(api_key=api_key)
        model = os.getenv("OPENAI_IMAGE_MODEL", "gpt-image-1").strip()
        size = os.getenv("OPENAI_IMAGE_SIZE", "1536x1024").strip()
        quality = os.getenv("OPENAI_IMAGE_QUALITY", "low").strip()
        resp = client.images.generate(
            model=model,
            prompt=f"Professional business presentation image: {query}. Clean composition, no text.",
            size=size,
            quality=quality,
        )
        if not resp.data:
            return None
        first = resp.data[0]
        b64 = getattr(first, "b64_json", None)
        if b64:
            destination.write_bytes(base64.b64decode(b64))
            return destination
        img_url = getattr(first, "url", None)
        if img_url:
            return _download_image(img_url, destination)
        return None
    except Exception:  
        return None




















# import json
# import io
# import os
# import time
# import base64
# from contextlib import redirect_stdout
# from pathlib import Path
# from typing import Any, Dict, List, Optional
# from urllib.parse import quote_plus
# from urllib.request import Request, urlopen
# from urllib.parse import urlparse

# from linkup import LinkupClient
# from pptx import Presentation
# from pptx.dml.color import RGBColor
# from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
# from pptx.util import Inches, Pt


# class MarketSearchTool:
#     """Wrapper around Linkup search with retry/backoff and bounded multi-query repair search."""

#     def __init__(self) -> None:
#         self._client = None
#         self._max_attempts = self._read_int_env("LINKUP_SEARCH_RETRIES", 2) + 1
#         self._retry_backoff_seconds = self._read_float_env("LINKUP_SEARCH_RETRY_BACKOFF_SECONDS", 0.8)
#         allowed_env = os.getenv("LINKUP_ALLOWED_DOMAINS", "").strip()
#         blocked_env = os.getenv("LINKUP_BLOCKED_DOMAINS", "").strip()
#         default_blocked = {
#             "github.com",
#             "marketreportsworld.com",
#             "verifiedmarketreports.com",
#             "futuredatastats.com",
#             "dataintelo.com",
#             "marketintelo.com",
#             "htfmarketinsights.com",
#         }
#         self._allowed_domains = {d.strip().lower() for d in allowed_env.split(",") if d.strip()}
#         blocked_domains = {d.strip().lower() for d in blocked_env.split(",") if d.strip()}
#         self._blocked_domains = blocked_domains or default_blocked
#         try:
#             self._client = LinkupClient()
#         except Exception:
#             self._client = None

#     @staticmethod
#     def _read_int_env(name: str, default: int) -> int:
#         try:
#             return max(0, int(os.getenv(name, str(default))))
#         except (TypeError, ValueError):
#             return default

#     @staticmethod
#     def _read_float_env(name: str, default: float) -> float:
#         try:
#             return max(0.0, float(os.getenv(name, str(default))))
#         except (TypeError, ValueError):
#             return default

#     @staticmethod
#     def _extract_domain(url: str) -> str:
#         try:
#             parsed = urlparse(url)
#             domain = parsed.netloc.lower()
#             if domain.startswith("www."):
#                 domain = domain[4:]
#             return domain
#         except Exception:
#             return ""

#     def _is_domain_allowed(self, domain: str) -> bool:
#         if not domain:
#             return False
#         if self._allowed_domains:
#             return domain in self._allowed_domains
#         return domain not in self._blocked_domains

#     @staticmethod
#     def source_quality_score(url: str) -> float:
#         url_l = (url or "").lower()
#         high_quality = [
#             ".gov", ".edu", "who.int", "nih.gov", "cms.gov", "hhs.gov",
#             "oecd.org", "worldbank.org", "mckinsey.com", "deloitte.com",
#             "pwc.com", "cbinsights.com", "rockhealth.com", "gartner.com",
#         ]
#         low_quality = ["medium.com", "substack.com", "wikipedia.org", "marketreportsworld.com", "verifiedmarketreports.com"]
#         if any(m in url_l for m in high_quality):
#             return 1.0
#         if any(m in url_l for m in low_quality):
#             return 0.35
#         return 0.65

#     def search(self, query: str) -> Dict[str, Any]:
#         if self._client is None:
#             return {
#                 "status": "error",
#                 "query": query,
#                 "results_json": "[]",
#                 "results": [],
#                 "sources": [],
#                 "dropped_sources": 0,
#                 "source_count_raw": 0,
#                 "attempts": 0,
#                 "retry_errors": [],
#                 "error": "Set LINKUP_API_KEY to enable live web search.",
#             }

#         retry_errors: List[str] = []
#         max_attempts = max(1, int(self._max_attempts))
#         for attempt in range(1, max_attempts + 1):
#             try:
#                 response = self._client.search(
#                     query=query,
#                     depth="standard",
#                     output_type="searchResults",
#                 )
#                 raw_results = getattr(response, "results", [])
#                 compact = []
#                 seen_urls = set()
#                 dropped_count = 0
#                 for r in raw_results:
#                     if isinstance(r, dict):
#                         title = r.get("name", "")
#                         url = r.get("url", "")
#                         content = r.get("content", "")
#                     else:
#                         title = getattr(r, "name", "")
#                         url = getattr(r, "url", "")
#                         content = getattr(r, "content", "")
#                     if not url or url in seen_urls:
#                         continue
#                     domain = self._extract_domain(url)
#                     if not self._is_domain_allowed(domain):
#                         dropped_count += 1
#                         continue
#                     seen_urls.add(url)
#                     compact.append({
#                         "title": title,
#                         "url": url,
#                         "domain": domain,
#                         "content": str(content)[:400],
#                         "quality_score": self.source_quality_score(url),
#                     })
#                 sources = [item["url"] for item in compact if item.get("url")]
#                 return {
#                     "status": "ok",
#                     "query": query,
#                     "results_json": json.dumps(compact, indent=2),
#                     "results": compact,
#                     "sources": sources,
#                     "dropped_sources": dropped_count,
#                     "source_count_raw": len(sources),
#                     "attempts": attempt,
#                     "retry_errors": retry_errors,
#                     "error": "",
#                 }
#             except Exception as exc:
#                 retry_errors.append(f"attempt_{attempt}: {exc}")
#                 if attempt < max_attempts and self._retry_backoff_seconds > 0:
#                     time.sleep(self._retry_backoff_seconds * attempt)

#         return {
#             "status": "error",
#             "query": query,
#             "results_json": "[]",
#             "results": [],
#             "sources": [],
#             "dropped_sources": 0,
#             "source_count_raw": 0,
#             "attempts": max_attempts,
#             "retry_errors": retry_errors,
#             "error": "Market search failed after " f"{max_attempts} attempt(s): {'; '.join(retry_errors)}",
#         }

#     def search_many(self, queries: List[str], max_results_per_query: int = 3) -> Dict[str, Any]:
#         """Run bounded per-claim searches and deduplicate results by URL."""
#         merged: List[Dict[str, Any]] = []
#         seen_urls = set()
#         statuses: List[Dict[str, Any]] = []
#         dropped_sources = 0
#         errors: List[str] = []
#         for query in queries:
#             clean_query = str(query or "").strip()
#             if not clean_query:
#                 continue
#             payload = self.search(clean_query)
#             statuses.append({
#                 "query": clean_query,
#                 "status": payload.get("status", "unknown"),
#                 "source_count": len(payload.get("sources", []) or []),
#                 "attempts": payload.get("attempts", 0),
#                 "retry_errors": payload.get("retry_errors", []),
#                 "error": payload.get("error", ""),
#             })
#             dropped_sources += int(payload.get("dropped_sources", 0) or 0)
#             if payload.get("status") != "ok":
#                 if payload.get("error"):
#                     errors.append(str(payload.get("error")))
#                 continue
#             for item in (payload.get("results", []) or [])[:max_results_per_query]:
#                 if not isinstance(item, dict):
#                     continue
#                 url = str(item.get("url", "") or "")
#                 if url and url in seen_urls:
#                     continue
#                 if url:
#                     seen_urls.add(url)
#                 item = dict(item)
#                 item["quality_score"] = self.source_quality_score(url)
#                 merged.append(item)
#         merged = sorted(merged, key=lambda item: float(item.get("quality_score", 0.0) or 0.0), reverse=True)
#         return {
#             "status": "ok" if merged else "error",
#             "query": " | ".join(str(q).strip() for q in queries if str(q).strip()),
#             "queries": [str(q).strip() for q in queries if str(q).strip()],
#             "query_statuses": statuses,
#             "results": merged,
#             "results_json": json.dumps(merged, indent=2),
#             "sources": [item["url"] for item in merged if item.get("url")],
#             "source_count_raw": len(merged),
#             "dropped_sources": dropped_sources,
#             "error": "; ".join(errors) if errors else ("" if merged else "No search results returned."),
#         }


# class BusinessCalcTool:
#     """Python REPL helper for simple financial projections."""

#     def run(self, code: str) -> str:
#         try:
#             buf = io.StringIO()
#             globals_dict = {"__builtins__": {"print": print, "range": range, "len": len}}
#             with redirect_stdout(buf):
#                 exec(code, globals_dict, {})
#             return buf.getvalue().strip()
#         except Exception as exc:  
#             return f"Python calc failed: {exc}"


# class GoogleTrendsTool:
#     """Google Trends wrapper via pytrends."""

#     def __init__(self) -> None:
#         self.hl = os.getenv("GOOGLE_TRENDS_HL", "en-US").strip()
#         self.tz = int(os.getenv("GOOGLE_TRENDS_TZ", "360").strip() or "360")
#         self.geo = os.getenv("GOOGLE_TRENDS_GEO", "US").strip()
#         self.timeframe = os.getenv("GOOGLE_TRENDS_TIMEFRAME", "today 12-m").strip()
#         self.connect_timeout = int(os.getenv("GOOGLE_TRENDS_CONNECT_TIMEOUT", "6").strip())
#         self.read_timeout = int(os.getenv("GOOGLE_TRENDS_READ_TIMEOUT", "12").strip())
#         self.retries = int(os.getenv("GOOGLE_TRENDS_RETRIES", "2").strip())
#         self.backoff_factor = float(os.getenv("GOOGLE_TRENDS_BACKOFF_FACTOR", "0.4").strip())
#         self.max_keywords = int(os.getenv("GOOGLE_TRENDS_MAX_KEYWORDS", "5").strip())
#         self._init_error = ""
#         try:
#             from pytrends.request import TrendReq

#             self._client = TrendReq(
#                 hl=self.hl,
#                 tz=self.tz,
#                 timeout=(self.connect_timeout, self.read_timeout),
#                 retries=self.retries,
#                 backoff_factor=self.backoff_factor,
#             )
#         except Exception as exc:  
#             self._client = None
#             self._init_error = str(exc)

#     def fetch(self, keywords: List[str]) -> Dict[str, Any]:
#         clean_keywords = [k.strip() for k in keywords if k and k.strip()][: self.max_keywords]
#         if not clean_keywords:
#             return {
#                 "status": "error",
#                 "keywords": [],
#                 "data": {},
#                 "error": "No keywords provided for trends lookup.",
#             }
#         if self._client is None:
#             return {
#                 "status": "error",
#                 "keywords": clean_keywords,
#                 "data": {},
#                 "error": (
#                     "Google Trends unavailable. Ensure pytrends is installed in the same "
#                     f"Python environment. Details: {self._init_error or 'import failed'}"
#                 ),
#             }

#         per_keyword: Dict[str, Any] = {}
#         failures: List[str] = []
#         for kw in clean_keywords:
#             attempts = max(1, self.retries + 1)
#             last_error = ""
#             for attempt in range(attempts):
#                 try:
#                     self._client.build_payload([kw], timeframe=self.timeframe, geo=self.geo)
#                     iot = self._client.interest_over_time()
#                     related = self._client.related_queries() or {}

#                     summary = {
#                         "average_interest": 0.0,
#                         "latest_interest": 0,
#                         "peak_interest": 0,
#                         "top_queries": [],
#                         "rising_queries": [],
#                     }
#                     if not iot.empty and kw in iot.columns:
#                         values = [int(v) for v in iot[kw].tolist()]
#                         if values:
#                             summary["average_interest"] = round(sum(values) / len(values), 2)
#                             summary["latest_interest"] = values[-1]
#                             summary["peak_interest"] = max(values)

#                     kw_related = related.get(kw, {})
#                     top_df = kw_related.get("top")
#                     if top_df is not None and not top_df.empty and "query" in top_df.columns:
#                         summary["top_queries"] = [
#                             str(q) for q in top_df["query"].head(5).tolist()
#                         ]
#                     rising_df = kw_related.get("rising")
#                     if (
#                         rising_df is not None
#                         and not rising_df.empty
#                         and "query" in rising_df.columns
#                     ):
#                         summary["rising_queries"] = [
#                             str(q) for q in rising_df["query"].head(5).tolist()
#                         ]
#                     per_keyword[kw] = summary
#                     last_error = ""
#                     break
#                 except Exception as exc:  
#                     last_error = str(exc)
#                     if attempt < attempts - 1:
#                         sleep_s = self.backoff_factor * (2**attempt)
#                         time.sleep(max(0.2, sleep_s))
#             if last_error:
#                 failures.append(f"{kw}: {last_error}")

#         if not per_keyword:
#             return {
#                 "status": "error",
#                 "keywords": clean_keywords,
#                 "data": {},
#                 "error": "; ".join(failures)[:500] or "Google Trends request failed.",
#             }

#         status = "ok" if not failures else "partial"
#         return {
#             "status": status,
#             "keywords": clean_keywords,
#             "data": per_keyword,
#             "error": "; ".join(failures)[:500] if failures else "",
#         }


# class ScenarioAnalysisTool:
#     """Deterministic scenario projections for startup plans."""

#     def run(
#         self,
#         users_year1: int = 8000,
#         arpu_monthly: float = 12.0,
#         gross_margin: float = 0.72,
#     ) -> Dict[str, Any]:
#         scenarios = {
#             "conservative": {"users_multiplier": 0.7, "arpu_multiplier": 0.9},
#             "base": {"users_multiplier": 1.0, "arpu_multiplier": 1.0},
#             "aggressive": {"users_multiplier": 1.35, "arpu_multiplier": 1.15},
#         }
#         output: Dict[str, Any] = {}
#         for name, mult in scenarios.items():
#             users = int(round(users_year1 * mult["users_multiplier"]))
#             arpu = arpu_monthly * mult["arpu_multiplier"]
#             revenue = users * arpu * 12
#             gross_profit = revenue * gross_margin
#             output[name] = {
#                 "users": users,
#                 "arpu_monthly": round(arpu, 2),
#                 "annual_revenue": round(revenue, 2),
#                 "gross_profit": round(gross_profit, 2),
#                 "gross_margin": gross_margin,
#             }
#         return output


# def generate_pitch_deck(slides: Dict[str, str], output_path: str) -> str:
#     """Create a styled .pptx deck from text content with optional visual enrichment."""
#     path = Path(output_path)
#     path.parent.mkdir(parents=True, exist_ok=True)
#     image_dir = path.parent / "images"
#     image_dir.mkdir(parents=True, exist_ok=True)

#     prs = Presentation()
#     prs.slide_width = Inches(13.333)
#     prs.slide_height = Inches(7.5)

#     title_layout = prs.slide_layouts[0]
#     body_layout = prs.slide_layouts[1]

#     title_slide = prs.slides.add_slide(title_layout)
#     _apply_background(title_slide, RGBColor(17, 24, 39), RGBColor(31, 41, 55))
#     title_slide.shapes.title.text = slides.get("title", "Startup Pitch")
#     title_slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(44)
#     title_slide.shapes.title.text_frame.paragraphs[0].font.bold = True
#     title_slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(245, 245, 245)
#     title_slide.placeholders[1].text = slides.get("subtitle", "AI Startup Pitch Refinery")
#     subtitle_p = title_slide.placeholders[1].text_frame.paragraphs[0]
#     subtitle_p.font.size = Pt(22)
#     subtitle_p.font.color.rgb = RGBColor(209, 213, 219)

#     # Accent stripe for title slide
#     stripe = title_slide.shapes.add_shape(
#         MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(6.8), Inches(13.333), Inches(0.7)
#     )
#     stripe.fill.solid()
#     stripe.fill.fore_color.rgb = RGBColor(59, 130, 246)
#     stripe.line.fill.background()

#     ordered_sections: List[str] = [
#         "problem",
#         "solution",
#         "market",
#         "business_model",
#         "competitive_advantage",
#         "financials",
#     ]

#     headers = {
#         "problem": "Problem",
#         "solution": "Solution",
#         "market": "Market",
#         "business_model": "Business Model",
#         "competitive_advantage": "Competitive Advantage",
#         "financials": "Financials",
#     }

#     for idx, key in enumerate(ordered_sections):
#         slide = prs.slides.add_slide(body_layout)
#         if len(slide.placeholders) > 1:
#             ph = slide.placeholders[1]
#             ph_el = ph._element
#             ph_el.getparent().remove(ph_el)
#         section_text = slides.get(key, "TBD")
#         accent_color = _section_color(idx)
#         _apply_background(slide, RGBColor(249, 250, 251), RGBColor(243, 244, 246))
#         _style_section_title(slide, headers[key], accent_color)

#         content_box = slide.shapes.add_shape(
#             MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.5), Inches(7.3), Inches(5.3)
#         )
#         content_box.fill.solid()
#         content_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
#         content_box.fill.transparency = 0.08
#         content_box.line.color.rgb = RGBColor(229, 231, 235)
#         content_box.line.width = Pt(1.5)

#         _fill_content_text(content_box.text_frame, section_text)

#         image_query = f"{headers[key]} startup business"
#         image_path = _fetch_slide_image(image_query, image_dir / f"{key}.jpg")

#         image_left = Inches(8.3)
#         image_top = Inches(1.55)
#         image_width = Inches(4.2)
#         image_height = Inches(5.2)
#         if image_path:
#             _add_picture_contain(
#                 slide,
#                 str(image_path),
#                 image_left,
#                 image_top,
#                 image_width,
#                 image_height,
#             )
#         else:
#             fallback = slide.shapes.add_shape(
#                 MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, image_left, image_top, image_width, image_height
#             )
#             fallback.fill.solid()
#             fallback.fill.fore_color.rgb = RGBColor(224, 231, 255)
#             fallback.line.color.rgb = accent_color
#             fallback_text = fallback.text_frame
#             fallback_text.text = "Visual Placeholder"
#             fallback_text.paragraphs[0].font.bold = True
#             fallback_text.paragraphs[0].font.size = Pt(18)
#             fallback_text.paragraphs[0].font.color.rgb = RGBColor(55, 65, 81)

#     prs.save(path)
#     return str(path)


# def _section_color(idx: int) -> RGBColor:
#     palette = [
#         RGBColor(37, 99, 235),
#         RGBColor(16, 185, 129),
#         RGBColor(249, 115, 22),
#         RGBColor(168, 85, 247),
#         RGBColor(244, 63, 94),
#         RGBColor(14, 165, 233),
#     ]
#     return palette[idx % len(palette)]


# def _apply_background(slide, top: RGBColor, bottom: RGBColor) -> None:
#     fill = slide.background.fill
#     fill.solid()
#     fill.fore_color.rgb = top


# def _style_section_title(slide, title: str, accent: RGBColor) -> None:
#     title_shape = slide.shapes.title
#     title_shape.text = title
#     p = title_shape.text_frame.paragraphs[0]
#     p.font.size = Pt(36)
#     p.font.bold = True
#     p.font.color.rgb = RGBColor(17, 24, 39)

#     accent_bar = slide.shapes.add_shape(
#         MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.7), Inches(1.2), Inches(2.2), Inches(0.08)
#     )
#     accent_bar.fill.solid()
#     accent_bar.fill.fore_color.rgb = accent
#     accent_bar.line.fill.background()


# def _fill_content_text(text_frame, content: str) -> None:
#     text_frame.clear()
#     lines = [ln.strip() for ln in content.split("\n") if ln.strip()]
#     if not lines:
#         lines = ["TBD"]
#     first = text_frame.paragraphs[0]
#     first.text = lines[0]
#     first.font.size = Pt(20)
#     first.font.bold = True
#     first.font.color.rgb = RGBColor(17, 24, 39)
#     for line in lines[1:]:
#         p = text_frame.add_paragraph()
#         p.text = line
#         p.level = 1
#         p.font.size = Pt(16)
#         p.font.color.rgb = RGBColor(55, 65, 81)


# def _add_picture_contain(slide, image_path: str, left, top, width, height) -> None:
#     """Place image inside target box without distortion and centered."""
#     pic = slide.shapes.add_picture(image_path, left, top)
#     iw, ih = float(pic.width), float(pic.height)
#     bw, bh = float(width), float(height)
#     if iw <= 0 or ih <= 0:
#         return

#     scale = min(bw / iw, bh / ih)
#     new_w = int(iw * scale)
#     new_h = int(ih * scale)
#     pic.width = new_w
#     pic.height = new_h
#     pic.left = int(float(left) + (bw - new_w) / 2)
#     pic.top = int(float(top) + (bh - new_h) / 2)


# def _fetch_slide_image(query: str, destination: Path) -> Optional[Path]:
#     reuse_cache = os.getenv("IMAGE_REUSE_CACHE", "1").strip().lower() in {"1", "true", "yes", "on"}
#     fetch_enabled = os.getenv("IMAGE_FETCH_ENABLED", "1").strip().lower() in {
#         "1",
#         "true",
#         "yes",
#         "on",
#     }
#     if reuse_cache and destination.exists():
#         return destination
#     if not fetch_enabled:
#         return destination if destination.exists() else None

#     provider = os.getenv("IMAGE_PROVIDER", "auto").strip().lower()
#     providers = [provider]
#     if provider == "auto":
#         providers = ["unsplash", "openai"]

#     for p in providers:
#         if p == "openai":
#             got = _fetch_from_openai_image(query, destination)
#         else:
#             got = _fetch_from_unsplash(query, destination)
#         if got:
#             return got
#     return None


# def _download_image(url: str, destination: Path, headers: Optional[Dict[str, str]] = None) -> Optional[Path]:
#     try:
#         req = Request(url, headers=headers or {"User-Agent": "Mozilla/5.0"})
#         with urlopen(req, timeout=20) as resp:  # noqa: S310
#             content_type = str(resp.headers.get("Content-Type", "")).lower()
#             if "image" not in content_type:
#                 return None
#             data = resp.read()
#         if not data:
#             return None
#         destination.write_bytes(data)
#         return destination
#     except Exception:  
#         return None


# def _fetch_from_unsplash(query: str, destination: Path) -> Optional[Path]:
#     url = f"https://source.unsplash.com/1600x900/?{quote_plus(query)}"
#     return _download_image(url, destination)


# def _fetch_from_openai_image(query: str, destination: Path) -> Optional[Path]:
#     enabled = os.getenv("OPENAI_IMAGE_ENABLED", "0").strip().lower() in {"1", "true", "yes", "on"}
#     if not enabled:
#         return None
#     api_key = os.getenv("OPENAI_API_KEY", "").strip()
#     if not api_key:
#         return None

#     try:
#         from openai import OpenAI

#         client = OpenAI(api_key=api_key)
#         model = os.getenv("OPENAI_IMAGE_MODEL", "gpt-image-1").strip()
#         size = os.getenv("OPENAI_IMAGE_SIZE", "1536x1024").strip()
#         quality = os.getenv("OPENAI_IMAGE_QUALITY", "low").strip()
#         resp = client.images.generate(
#             model=model,
#             prompt=f"Professional business presentation image: {query}. Clean composition, no text.",
#             size=size,
#             quality=quality,
#         )
#         if not resp.data:
#             return None
#         first = resp.data[0]
#         b64 = getattr(first, "b64_json", None)
#         if b64:
#             destination.write_bytes(base64.b64decode(b64))
#             return destination
#         img_url = getattr(first, "url", None)
#         if img_url:
#             return _download_image(img_url, destination)
#         return None
#     except Exception:  
#         return None






















# # import json
# # import io
# # import os
# # import time
# # import base64
# # from contextlib import redirect_stdout
# # from pathlib import Path
# # from typing import Any, Dict, List, Optional
# # from urllib.parse import quote_plus
# # from urllib.request import Request, urlopen
# # from urllib.parse import urlparse

# # from linkup import LinkupClient
# # from pptx import Presentation
# # from pptx.dml.color import RGBColor
# # from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
# # from pptx.util import Inches, Pt


# # class MarketSearchTool:
# #     """Wrapper around Linkup search with graceful fallback when key is missing."""

# #     def __init__(self) -> None:
# #         self._client = None
# #         self._max_attempts = self._read_int_env("LINKUP_SEARCH_RETRIES", 2) + 1
# #         self._retry_backoff_seconds = self._read_float_env(
# #             "LINKUP_SEARCH_RETRY_BACKOFF_SECONDS",
# #             0.8,
# #         )
# #         allowed_env = os.getenv("LINKUP_ALLOWED_DOMAINS", "").strip()
# #         blocked_env = os.getenv("LINKUP_BLOCKED_DOMAINS", "").strip()
# #         default_blocked = {
# #             "github.com",
# #             "marketreportsworld.com",
# #             "verifiedmarketreports.com",
# #             "futuredatastats.com",
# #             "dataintelo.com",
# #             "marketintelo.com",
# #             "htfmarketinsights.com",
# #         }
# #         self._allowed_domains = {
# #             d.strip().lower() for d in allowed_env.split(",") if d.strip()
# #         }
# #         blocked_domains = {d.strip().lower() for d in blocked_env.split(",") if d.strip()}
# #         self._blocked_domains = blocked_domains or default_blocked
# #         try:
# #             self._client = LinkupClient()
# #         except Exception:  
# #             self._client = None

# #     @staticmethod
# #     def _read_int_env(name: str, default: int) -> int:
# #         try:
# #             return max(0, int(os.getenv(name, str(default))))
# #         except (TypeError, ValueError):
# #             return default

# #     @staticmethod
# #     def _read_float_env(name: str, default: float) -> float:
# #         try:
# #             return max(0.0, float(os.getenv(name, str(default))))
# #         except (TypeError, ValueError):
# #             return default

# #     @staticmethod
# #     def _extract_domain(url: str) -> str:
# #         try:
# #             parsed = urlparse(url)
# #             domain = parsed.netloc.lower()
# #             if domain.startswith("www."):
# #                 domain = domain[4:]
# #             return domain
# #         except Exception:  
# #             return ""

# #     def _is_domain_allowed(self, domain: str) -> bool:
# #         if not domain:
# #             return False
# #         if self._allowed_domains:
# #             return domain in self._allowed_domains
# #         return domain not in self._blocked_domains

# #     @staticmethod
# #     def source_quality_score(url: str) -> float:
# #         """Simple source-quality prior for claim-level repair evidence."""
# #         url_l = (url or "").lower()
# #         high_quality_markers = [
# #             ".gov",
# #             ".edu",
# #             "who.int",
# #             "nih.gov",
# #             "cms.gov",
# #             "hhs.gov",
# #             "oecd.org",
# #             "worldbank.org",
# #             "mckinsey.com",
# #             "gartner.com",
# #             "deloitte.com",
# #             "pwc.com",
# #             "cbinsights.com",
# #             "rockhealth.com",
# #         ]
# #         low_quality_markers = [
# #             "medium.com",
# #             "substack.com",
# #             "wikipedia.org",
# #             "marketreportsworld.com",
# #             "verifiedmarketreports.com",
# #         ]
# #         if any(marker in url_l for marker in high_quality_markers):
# #             return 1.0
# #         if any(marker in url_l for marker in low_quality_markers):
# #             return 0.35
# #         return 0.65

# #     def search(self, query: str) -> Dict[str, Any]:
# #         if self._client is None:
# #             return {
# #                 "status": "error",
# #                 "query": query,
# #                 "results_json": "[]",
# #                 "results": [],
# #                 "sources": [],
# #                 "dropped_sources": 0,
# #                 "attempts": 0,
# #                 "retry_errors": [],
# #                 "error": "Set LINKUP_API_KEY to enable live web search.",
# #             }

# #         retry_errors: List[str] = []
# #         max_attempts = max(1, int(self._max_attempts))
# #         for attempt in range(1, max_attempts + 1):
# #             try:
# #                 response = self._client.search(
# #                     query=query,
# #                     depth="standard",
# #                     output_type="searchResults",
# #                 )
# #                 raw_results = getattr(response, "results", [])
# #                 compact = []
# #                 seen_urls = set()
# #                 dropped_count = 0
# #                 for r in raw_results:
# #                     if isinstance(r, dict):
# #                         title = r.get("name", "")
# #                         url = r.get("url", "")
# #                         content = r.get("content", "")
# #                     else:
# #                         title = getattr(r, "name", "")
# #                         url = getattr(r, "url", "")
# #                         content = getattr(r, "content", "")

# #                     if not url or url in seen_urls:
# #                         continue
# #                     domain = self._extract_domain(url)
# #                     if not self._is_domain_allowed(domain):
# #                         dropped_count += 1
# #                         continue
# #                     seen_urls.add(url)
# #                     compact.append(
# #                         {
# #                             "title": title,
# #                             "url": url,
# #                             "domain": domain,
# #                             "content": str(content)[:400],
# #                             "quality_score": self.source_quality_score(url),
# #                         }
# #                     )
# #                 sources = [item["url"] for item in compact if item.get("url")]
# #                 return {
# #                     "status": "ok",
# #                     "query": query,
# #                     "results_json": json.dumps(compact, indent=2),
# #                     "results": compact,
# #                     "sources": sources,
# #                     "dropped_sources": dropped_count,
# #                     "attempts": attempt,
# #                     "retry_errors": retry_errors,
# #                     "error": "",
# #                 }
# #             except Exception as exc:
# #                 retry_errors.append(f"attempt_{attempt}: {exc}")
# #                 if attempt < max_attempts and self._retry_backoff_seconds > 0:
# #                     time.sleep(self._retry_backoff_seconds * attempt)

# #         return {
# #             "status": "error",
# #             "query": query,
# #             "results_json": "[]",
# #             "results": [],
# #             "sources": [],
# #             "dropped_sources": 0,
# #             "attempts": max_attempts,
# #             "retry_errors": retry_errors,
# #             "error": "Market search failed after "
# #             f"{max_attempts} attempt(s): {'; '.join(retry_errors)}",
# #         }

# #     def search_many(
# #         self,
# #         queries: List[str],
# #         max_results_per_query: int = 3,
# #     ) -> Dict[str, Any]:
# #         """Run bounded per-claim searches and deduplicate by URL."""
# #         merged: List[Dict[str, Any]] = []
# #         seen_urls = set()
# #         statuses: List[Dict[str, Any]] = []
# #         dropped_sources = 0
# #         errors: List[str] = []
# #         for query in queries:
# #             clean_query = str(query or "").strip()
# #             if not clean_query:
# #                 continue
# #             payload = self.search(clean_query)
# #             statuses.append(
# #                 {
# #                     "query": clean_query,
# #                     "status": payload.get("status", "unknown"),
# #                     "source_count": len(payload.get("sources", []) or []),
# #                     "attempts": payload.get("attempts", 0),
# #                     "retry_errors": payload.get("retry_errors", []),
# #                     "error": payload.get("error", ""),
# #                 }
# #             )
# #             dropped_sources += int(payload.get("dropped_sources", 0) or 0)
# #             if payload.get("status") != "ok":
# #                 if payload.get("error"):
# #                     errors.append(str(payload.get("error")))
# #                 continue
# #             for item in (payload.get("results", []) or [])[:max_results_per_query]:
# #                 if not isinstance(item, dict):
# #                     continue
# #                 url = str(item.get("url", "") or "")
# #                 if url and url in seen_urls:
# #                     continue
# #                 if url:
# #                     seen_urls.add(url)
# #                 item = dict(item)
# #                 item["quality_score"] = self.source_quality_score(url)
# #                 merged.append(item)
# #         merged = sorted(
# #             merged,
# #             key=lambda item: float(item.get("quality_score", 0.0) or 0.0),
# #             reverse=True,
# #         )
# #         return {
# #             "status": "ok" if merged else "error",
# #             "query": " | ".join(str(q).strip() for q in queries if str(q).strip()),
# #             "queries": [str(q).strip() for q in queries if str(q).strip()],
# #             "query_statuses": statuses,
# #             "results": merged,
# #             "results_json": json.dumps(merged, indent=2),
# #             "sources": [item["url"] for item in merged if item.get("url")],
# #             "dropped_sources": dropped_sources,
# #             "error": "; ".join(errors) if errors else ("" if merged else "No search results returned."),
# #         }


# # class BusinessCalcTool:
# #     """Python REPL helper for simple financial projections."""

# #     def run(self, code: str) -> str:
# #         try:
# #             buf = io.StringIO()
# #             globals_dict = {"__builtins__": {"print": print, "range": range, "len": len}}
# #             with redirect_stdout(buf):
# #                 exec(code, globals_dict, {})
# #             return buf.getvalue().strip()
# #         except Exception as exc:  
# #             return f"Python calc failed: {exc}"


# # class GoogleTrendsTool:
# #     """Google Trends wrapper via pytrends."""

# #     def __init__(self) -> None:
# #         self.hl = os.getenv("GOOGLE_TRENDS_HL", "en-US").strip()
# #         self.tz = int(os.getenv("GOOGLE_TRENDS_TZ", "360").strip() or "360")
# #         self.geo = os.getenv("GOOGLE_TRENDS_GEO", "US").strip()
# #         self.timeframe = os.getenv("GOOGLE_TRENDS_TIMEFRAME", "today 12-m").strip()
# #         self.connect_timeout = int(os.getenv("GOOGLE_TRENDS_CONNECT_TIMEOUT", "6").strip())
# #         self.read_timeout = int(os.getenv("GOOGLE_TRENDS_READ_TIMEOUT", "12").strip())
# #         self.retries = int(os.getenv("GOOGLE_TRENDS_RETRIES", "2").strip())
# #         self.backoff_factor = float(os.getenv("GOOGLE_TRENDS_BACKOFF_FACTOR", "0.4").strip())
# #         self.max_keywords = int(os.getenv("GOOGLE_TRENDS_MAX_KEYWORDS", "5").strip())
# #         self._init_error = ""
# #         try:
# #             from pytrends.request import TrendReq

# #             self._client = TrendReq(
# #                 hl=self.hl,
# #                 tz=self.tz,
# #                 timeout=(self.connect_timeout, self.read_timeout),
# #                 retries=self.retries,
# #                 backoff_factor=self.backoff_factor,
# #             )
# #         except Exception as exc:  
# #             self._client = None
# #             self._init_error = str(exc)

# #     def fetch(self, keywords: List[str]) -> Dict[str, Any]:
# #         clean_keywords = [k.strip() for k in keywords if k and k.strip()][: self.max_keywords]
# #         if not clean_keywords:
# #             return {
# #                 "status": "error",
# #                 "keywords": [],
# #                 "data": {},
# #                 "error": "No keywords provided for trends lookup.",
# #             }
# #         if self._client is None:
# #             return {
# #                 "status": "error",
# #                 "keywords": clean_keywords,
# #                 "data": {},
# #                 "error": (
# #                     "Google Trends unavailable. Ensure pytrends is installed in the same "
# #                     f"Python environment. Details: {self._init_error or 'import failed'}"
# #                 ),
# #             }

# #         per_keyword: Dict[str, Any] = {}
# #         failures: List[str] = []
# #         for kw in clean_keywords:
# #             attempts = max(1, self.retries + 1)
# #             last_error = ""
# #             for attempt in range(attempts):
# #                 try:
# #                     self._client.build_payload([kw], timeframe=self.timeframe, geo=self.geo)
# #                     iot = self._client.interest_over_time()
# #                     related = self._client.related_queries() or {}

# #                     summary = {
# #                         "average_interest": 0.0,
# #                         "latest_interest": 0,
# #                         "peak_interest": 0,
# #                         "top_queries": [],
# #                         "rising_queries": [],
# #                     }
# #                     if not iot.empty and kw in iot.columns:
# #                         values = [int(v) for v in iot[kw].tolist()]
# #                         if values:
# #                             summary["average_interest"] = round(sum(values) / len(values), 2)
# #                             summary["latest_interest"] = values[-1]
# #                             summary["peak_interest"] = max(values)

# #                     kw_related = related.get(kw, {})
# #                     top_df = kw_related.get("top")
# #                     if top_df is not None and not top_df.empty and "query" in top_df.columns:
# #                         summary["top_queries"] = [
# #                             str(q) for q in top_df["query"].head(5).tolist()
# #                         ]
# #                     rising_df = kw_related.get("rising")
# #                     if (
# #                         rising_df is not None
# #                         and not rising_df.empty
# #                         and "query" in rising_df.columns
# #                     ):
# #                         summary["rising_queries"] = [
# #                             str(q) for q in rising_df["query"].head(5).tolist()
# #                         ]
# #                     per_keyword[kw] = summary
# #                     last_error = ""
# #                     break
# #                 except Exception as exc:  
# #                     last_error = str(exc)
# #                     if attempt < attempts - 1:
# #                         sleep_s = self.backoff_factor * (2**attempt)
# #                         time.sleep(max(0.2, sleep_s))
# #             if last_error:
# #                 failures.append(f"{kw}: {last_error}")

# #         if not per_keyword:
# #             return {
# #                 "status": "error",
# #                 "keywords": clean_keywords,
# #                 "data": {},
# #                 "error": "; ".join(failures)[:500] or "Google Trends request failed.",
# #             }

# #         status = "ok" if not failures else "partial"
# #         return {
# #             "status": status,
# #             "keywords": clean_keywords,
# #             "data": per_keyword,
# #             "error": "; ".join(failures)[:500] if failures else "",
# #         }


# # class ScenarioAnalysisTool:
# #     """Deterministic scenario projections for startup plans."""

# #     def run(
# #         self,
# #         users_year1: int = 8000,
# #         arpu_monthly: float = 12.0,
# #         gross_margin: float = 0.72,
# #     ) -> Dict[str, Any]:
# #         scenarios = {
# #             "conservative": {"users_multiplier": 0.7, "arpu_multiplier": 0.9},
# #             "base": {"users_multiplier": 1.0, "arpu_multiplier": 1.0},
# #             "aggressive": {"users_multiplier": 1.35, "arpu_multiplier": 1.15},
# #         }
# #         output: Dict[str, Any] = {}
# #         for name, mult in scenarios.items():
# #             users = int(round(users_year1 * mult["users_multiplier"]))
# #             arpu = arpu_monthly * mult["arpu_multiplier"]
# #             revenue = users * arpu * 12
# #             gross_profit = revenue * gross_margin
# #             output[name] = {
# #                 "users": users,
# #                 "arpu_monthly": round(arpu, 2),
# #                 "annual_revenue": round(revenue, 2),
# #                 "gross_profit": round(gross_profit, 2),
# #                 "gross_margin": gross_margin,
# #             }
# #         return output


# # def generate_pitch_deck(slides: Dict[str, str], output_path: str) -> str:
# #     """Create a styled .pptx deck from text content with optional visual enrichment."""
# #     path = Path(output_path)
# #     path.parent.mkdir(parents=True, exist_ok=True)
# #     image_dir = path.parent / "images"
# #     image_dir.mkdir(parents=True, exist_ok=True)

# #     prs = Presentation()
# #     prs.slide_width = Inches(13.333)
# #     prs.slide_height = Inches(7.5)

# #     title_layout = prs.slide_layouts[0]
# #     body_layout = prs.slide_layouts[1]

# #     title_slide = prs.slides.add_slide(title_layout)
# #     _apply_background(title_slide, RGBColor(17, 24, 39), RGBColor(31, 41, 55))
# #     title_slide.shapes.title.text = slides.get("title", "Startup Pitch")
# #     title_slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(44)
# #     title_slide.shapes.title.text_frame.paragraphs[0].font.bold = True
# #     title_slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(245, 245, 245)
# #     title_slide.placeholders[1].text = slides.get("subtitle", "AI Startup Pitch Refinery")
# #     subtitle_p = title_slide.placeholders[1].text_frame.paragraphs[0]
# #     subtitle_p.font.size = Pt(22)
# #     subtitle_p.font.color.rgb = RGBColor(209, 213, 219)

# #     # Accent stripe for title slide
# #     stripe = title_slide.shapes.add_shape(
# #         MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(6.8), Inches(13.333), Inches(0.7)
# #     )
# #     stripe.fill.solid()
# #     stripe.fill.fore_color.rgb = RGBColor(59, 130, 246)
# #     stripe.line.fill.background()

# #     ordered_sections: List[str] = [
# #         "problem",
# #         "solution",
# #         "market",
# #         "business_model",
# #         "competitive_advantage",
# #         "financials",
# #     ]

# #     headers = {
# #         "problem": "Problem",
# #         "solution": "Solution",
# #         "market": "Market",
# #         "business_model": "Business Model",
# #         "competitive_advantage": "Competitive Advantage",
# #         "financials": "Financials",
# #     }

# #     for idx, key in enumerate(ordered_sections):
# #         slide = prs.slides.add_slide(body_layout)
# #         if len(slide.placeholders) > 1:
# #             ph = slide.placeholders[1]
# #             ph_el = ph._element
# #             ph_el.getparent().remove(ph_el)
# #         section_text = slides.get(key, "TBD")
# #         accent_color = _section_color(idx)
# #         _apply_background(slide, RGBColor(249, 250, 251), RGBColor(243, 244, 246))
# #         _style_section_title(slide, headers[key], accent_color)

# #         content_box = slide.shapes.add_shape(
# #             MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.5), Inches(7.3), Inches(5.3)
# #         )
# #         content_box.fill.solid()
# #         content_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
# #         content_box.fill.transparency = 0.08
# #         content_box.line.color.rgb = RGBColor(229, 231, 235)
# #         content_box.line.width = Pt(1.5)

# #         _fill_content_text(content_box.text_frame, section_text)

# #         image_query = f"{headers[key]} startup business"
# #         image_path = _fetch_slide_image(image_query, image_dir / f"{key}.jpg")

# #         image_left = Inches(8.3)
# #         image_top = Inches(1.55)
# #         image_width = Inches(4.2)
# #         image_height = Inches(5.2)
# #         if image_path:
# #             _add_picture_contain(
# #                 slide,
# #                 str(image_path),
# #                 image_left,
# #                 image_top,
# #                 image_width,
# #                 image_height,
# #             )
# #         else:
# #             fallback = slide.shapes.add_shape(
# #                 MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, image_left, image_top, image_width, image_height
# #             )
# #             fallback.fill.solid()
# #             fallback.fill.fore_color.rgb = RGBColor(224, 231, 255)
# #             fallback.line.color.rgb = accent_color
# #             fallback_text = fallback.text_frame
# #             fallback_text.text = "Visual Placeholder"
# #             fallback_text.paragraphs[0].font.bold = True
# #             fallback_text.paragraphs[0].font.size = Pt(18)
# #             fallback_text.paragraphs[0].font.color.rgb = RGBColor(55, 65, 81)

# #     prs.save(path)
# #     return str(path)


# # def _section_color(idx: int) -> RGBColor:
# #     palette = [
# #         RGBColor(37, 99, 235),
# #         RGBColor(16, 185, 129),
# #         RGBColor(249, 115, 22),
# #         RGBColor(168, 85, 247),
# #         RGBColor(244, 63, 94),
# #         RGBColor(14, 165, 233),
# #     ]
# #     return palette[idx % len(palette)]


# # def _apply_background(slide, top: RGBColor, bottom: RGBColor) -> None:
# #     fill = slide.background.fill
# #     fill.solid()
# #     fill.fore_color.rgb = top


# # def _style_section_title(slide, title: str, accent: RGBColor) -> None:
# #     title_shape = slide.shapes.title
# #     title_shape.text = title
# #     p = title_shape.text_frame.paragraphs[0]
# #     p.font.size = Pt(36)
# #     p.font.bold = True
# #     p.font.color.rgb = RGBColor(17, 24, 39)

# #     accent_bar = slide.shapes.add_shape(
# #         MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.7), Inches(1.2), Inches(2.2), Inches(0.08)
# #     )
# #     accent_bar.fill.solid()
# #     accent_bar.fill.fore_color.rgb = accent
# #     accent_bar.line.fill.background()


# # def _fill_content_text(text_frame, content: str) -> None:
# #     text_frame.clear()
# #     lines = [ln.strip() for ln in content.split("\n") if ln.strip()]
# #     if not lines:
# #         lines = ["TBD"]
# #     first = text_frame.paragraphs[0]
# #     first.text = lines[0]
# #     first.font.size = Pt(20)
# #     first.font.bold = True
# #     first.font.color.rgb = RGBColor(17, 24, 39)
# #     for line in lines[1:]:
# #         p = text_frame.add_paragraph()
# #         p.text = line
# #         p.level = 1
# #         p.font.size = Pt(16)
# #         p.font.color.rgb = RGBColor(55, 65, 81)


# # def _add_picture_contain(slide, image_path: str, left, top, width, height) -> None:
# #     """Place image inside target box without distortion and centered."""
# #     pic = slide.shapes.add_picture(image_path, left, top)
# #     iw, ih = float(pic.width), float(pic.height)
# #     bw, bh = float(width), float(height)
# #     if iw <= 0 or ih <= 0:
# #         return

# #     scale = min(bw / iw, bh / ih)
# #     new_w = int(iw * scale)
# #     new_h = int(ih * scale)
# #     pic.width = new_w
# #     pic.height = new_h
# #     pic.left = int(float(left) + (bw - new_w) / 2)
# #     pic.top = int(float(top) + (bh - new_h) / 2)


# # def _fetch_slide_image(query: str, destination: Path) -> Optional[Path]:
# #     reuse_cache = os.getenv("IMAGE_REUSE_CACHE", "1").strip().lower() in {"1", "true", "yes", "on"}
# #     fetch_enabled = os.getenv("IMAGE_FETCH_ENABLED", "1").strip().lower() in {
# #         "1",
# #         "true",
# #         "yes",
# #         "on",
# #     }
# #     if reuse_cache and destination.exists():
# #         return destination
# #     if not fetch_enabled:
# #         return destination if destination.exists() else None

# #     provider = os.getenv("IMAGE_PROVIDER", "auto").strip().lower()
# #     providers = [provider]
# #     if provider == "auto":
# #         providers = ["unsplash", "openai"]

# #     for p in providers:
# #         if p == "openai":
# #             got = _fetch_from_openai_image(query, destination)
# #         else:
# #             got = _fetch_from_unsplash(query, destination)
# #         if got:
# #             return got
# #     return None


# # def _download_image(url: str, destination: Path, headers: Optional[Dict[str, str]] = None) -> Optional[Path]:
# #     try:
# #         req = Request(url, headers=headers or {"User-Agent": "Mozilla/5.0"})
# #         with urlopen(req, timeout=20) as resp:  # noqa: S310
# #             content_type = str(resp.headers.get("Content-Type", "")).lower()
# #             if "image" not in content_type:
# #                 return None
# #             data = resp.read()
# #         if not data:
# #             return None
# #         destination.write_bytes(data)
# #         return destination
# #     except Exception:  
# #         return None


# # def _fetch_from_unsplash(query: str, destination: Path) -> Optional[Path]:
# #     url = f"https://source.unsplash.com/1600x900/?{quote_plus(query)}"
# #     return _download_image(url, destination)


# # def _fetch_from_openai_image(query: str, destination: Path) -> Optional[Path]:
# #     enabled = os.getenv("OPENAI_IMAGE_ENABLED", "0").strip().lower() in {"1", "true", "yes", "on"}
# #     if not enabled:
# #         return None
# #     api_key = os.getenv("OPENAI_API_KEY", "").strip()
# #     if not api_key:
# #         return None

# #     try:
# #         from openai import OpenAI

# #         client = OpenAI(api_key=api_key)
# #         model = os.getenv("OPENAI_IMAGE_MODEL", "gpt-image-1").strip()
# #         size = os.getenv("OPENAI_IMAGE_SIZE", "1536x1024").strip()
# #         quality = os.getenv("OPENAI_IMAGE_QUALITY", "low").strip()
# #         resp = client.images.generate(
# #             model=model,
# #             prompt=f"Professional business presentation image: {query}. Clean composition, no text.",
# #             size=size,
# #             quality=quality,
# #         )
# #         if not resp.data:
# #             return None
# #         first = resp.data[0]
# #         b64 = getattr(first, "b64_json", None)
# #         if b64:
# #             destination.write_bytes(base64.b64decode(b64))
# #             return destination
# #         img_url = getattr(first, "url", None)
# #         if img_url:
# #             return _download_image(img_url, destination)
# #         return None
# #     except Exception:  
# #         return None
